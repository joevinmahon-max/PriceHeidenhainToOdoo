# ==========================================================
# PREPARATION PRIX HEIDENHAIN -> ODOO
# JMA
#
# - Detection automatique des colonnes
# - Colonnes Heidenhain sélectionnables
# - Colonnes par défaut définies automatiquement
# - Création des lignes SAV pour VG / PG
# - Conservation des prix PPC / SAV / HA
# - Correspondances Odoo automatiques
# - Colonnes Odoo existantes conservées dans leur ordre
# - Ajout des colonnes manquantes à l'emplacement choisi
# - Transfert des colonnes supplémentaires de l'étape 1
# - Conservation des ID Odoo existants
# - Index des références pour accélérer le traitement
# ==========================================================

import time
from io import BytesIO
from copy import copy

import openpyxl
import streamlit as st


# ==========================================================
# CONFIGURATION
# ==========================================================

st.set_page_config(
    page_title="Préparation Heidenhain vers Odoo",
    page_icon="📦",
    layout="wide",
)

st.title("📦 Préparation des prix HEIDENHAIN → Odoo")


# ==========================================================
# SESSION STATE
# ==========================================================

DEFAULT_STATE = {
    "heidenhain_result": None,
    "heidenhain_stats": None,
    "odoo_result": None,
    "odoo_stats": None,
    "heidenhain_processing": False,
    "odoo_processing": False,
}

for key, value in DEFAULT_STATE.items():
    if key not in st.session_state:
        st.session_state[key] = value


# ==========================================================
# CONSTANTES
# ==========================================================

STATUS_TO_DUPLICATE = {"VG", "PG"}

DEFAULT_HEIDENHAIN_COLUMNS = [
    "ID",
    "Description",
    "Marque",
    "Statut",
    "Groupe Produit",
    "Prix (PPC)",
    "Prix (SAV)",
    "Prix HA",
]

REQUIRED_HEIDENHAIN_COLUMNS = [
    "ID",
    "Description",
    "Marque",
    "Statut",
    "Groupe Produit",
    "Prix (PPC)",
    "Prix (SAV)",
    "Prix HA",
]


# ==========================================================
# CORRESPONDANCES HEIDENHAIN -> ODOO
# ==========================================================

ODOO_FIELD_CONFIGURATION = {
    "id": {
        "label": "ID Odoo",
        "source": None,
        "fixed": None,
    },
    "default_code": {
        "label": "Référence interne",
        "source": "ID",
        "fixed": None,
    },
    "name": {
        "label": "Nom",
        "source": "Description",
        "fixed": None,
    },
    "x_studio_marque_1": {
        "label": "Marque",
        "source": "Marque",
        "fixed": None,
    },
    "categ_id": {
        "label": "Catégorie",
        "source": "Groupe Produit",
        "fixed": None,
    },
    "x_studio_sales_status": {
        "label": "Sales Status",
        "source": "Statut",
        "fixed": None,
    },
    "list_price": {
        "label": "Prix de vente",
        "source": "__PRICE__",
        "fixed": None,
    },
    "seller_ids/partner_id": {
        "label": "Fournisseur",
        "source": None,
        "fixed": "HEIDENHAIN FRANCE",
    },
    "seller_ids/price": {
        "label": "Prix fournisseur",
        "source": "Prix HA",
        "fixed": None,
    },
    "purchase_ok": {
        "label": "Peut être acheté",
        "source": None,
        "fixed": "VRAI",
    },
    "sale_ok": {
        "label": "Peut être vendu",
        "source": None,
        "fixed": "VRAI",
    },
    "type": {
        "label": "Type de produit",
        "source": None,
        "fixed": "Consommable",
    },
    "invoice_policy": {
        "label": "Politique de facturation",
        "source": None,
        "fixed": "Quantités livrées",
    },
    "barcode": {
        "label": "Code-barres",
        "source": "__BARCODE__",
        "fixed": None,
    },
}

# Sources standards déjà associées à un champ Odoo.
SOURCE_TO_ODOO = {
    "ID": "default_code",
    "Description": "name",
    "Marque": "x_studio_marque_1",
    "Statut": "x_studio_sales_status",
    "Groupe Produit": "categ_id",
    "Prix (PPC)": "list_price",
    "Prix (SAV)": "list_price",
    "Prix HA": "seller_ids/price",
}

# ==========================================================
# UTILITAIRES
# ==========================================================

def normalize(value):
    if value is None:
        return ""
    return str(value).strip().upper()


def clean_reference(value):
    return normalize(value)


def is_sav(reference):
    return clean_reference(reference).endswith("_SAV")


def copy_style_safe(source, target):
    for attribute in (
        "font",
        "fill",
        "border",
        "alignment",
        "protection",
    ):
        try:
            setattr(
                target,
                attribute,
                copy(getattr(source, attribute)),
            )
        except Exception:
            pass

    try:
        target.number_format = source.number_format
    except Exception:
        pass


def get_sheet_names(uploaded_file):
    if uploaded_file is None:
        return []

    uploaded_file.seek(0)

    wb = openpyxl.load_workbook(
        uploaded_file,
        read_only=True,
        data_only=True,
    )

    names = list(wb.sheetnames)
    wb.close()

    uploaded_file.seek(0)
    return names


def get_headers(ws, header_row):
    headers = []
    seen = set()

    for cell in ws[header_row]:
        if cell.value in (None, ""):
            continue

        name = str(cell.value).strip()
        key = normalize(name)

        if not key or key in seen:
            continue

        seen.add(key)
        headers.append(name)

    return headers


def find_columns(ws, header_row, names=None):
    result = {}

    wanted = (
        {normalize(name): name for name in names}
        if names is not None
        else None
    )

    for cell in ws[header_row]:
        key = normalize(cell.value)

        if not key:
            continue

        if wanted is None:
            result.setdefault(str(cell.value).strip(), cell.column)
        elif key in wanted:
            result.setdefault(wanted[key], cell.column)

    return result


def detect_excel_headers(uploaded_file, sheet_name, header_row):
    if uploaded_file is None:
        return []

    uploaded_file.seek(0)

    wb = openpyxl.load_workbook(
        uploaded_file,
        read_only=True,
        data_only=True,
    )

    try:
        if sheet_name not in wb.sheetnames:
            return []

        return get_headers(wb[sheet_name], int(header_row))

    finally:
        wb.close()
        uploaded_file.seek(0)


# ==========================================================
# CONSTRUCTION DES COLONNES ODOO
# ==========================================================

def build_required_odoo_columns(
    heidenhain_output_columns,
    mapping_config,
):
    result = []
    seen = set()

    def add_column(name):
        key = normalize(name)
        if key and key not in seen:
            seen.add(key)
            result.append(name)

    # Ajouter les champs Odoo standards, sauf l'ID.
    # L'ID doit être conservé uniquement s'il existe déjà dans le modèle.
    for field in mapping_config:
        if normalize(field) == "id":
            continue
        add_column(field)

    # Ajouter uniquement les colonnes Heidenhain
    # qui ne sont pas déjà associées à un champ Odoo.
    mapped_sources = {
        normalize(config.get("source"))
        for config in mapping_config.values()
        if config.get("source")
        and not config["source"].startswith("__")
    }

    for name in heidenhain_output_columns:
        if normalize(name) in mapped_sources:
            continue

        if normalize(name) in {
            normalize("Prix (PPC)"),
            normalize("Prix (SAV)"),
        }:
            continue

        add_column(name)

    return result


# ==========================================================
# ETAPE 1 : TRAITEMENT HEIDENHAIN
# ==========================================================

def process_heidenhain(
    uploaded_file,
    sheet_name,
    header_row,
    data_start_row,
    status_column,
    id_column,
    output_columns,
    progress=None,
    status_display=None,
):
    if status_display:
        status_display.info("📂 Ouverture du fichier Heidenhain...")
    
    if not output_columns:
        raise ValueError("Sélectionne au moins une colonne.")

    for required in (id_column, status_column):
        if normalize(required) not in {
            normalize(x) for x in output_columns
        }:
            raise ValueError(
                f"La colonne {required} doit être conservée."
            )

    uploaded_file.seek(0)

    wb_formula = openpyxl.load_workbook(
        uploaded_file,
        data_only=False,
    )

    uploaded_file.seek(0)

    wb_values = openpyxl.load_workbook(
        uploaded_file,
        data_only=True,
    )

    try:
        if sheet_name not in wb_formula.sheetnames:
            raise ValueError(
                f"Feuille Heidenhain introuvable : {sheet_name}"
            )

        source_ws = wb_formula[sheet_name]
        values_ws = wb_values[sheet_name]

        source_columns = find_columns(
            source_ws,
            int(header_row),
            output_columns,
        )

        missing = [
            name
            for name in output_columns
            if name not in source_columns
        ]

        if missing:
            raise ValueError(
                "Colonnes Heidenhain introuvables : "
                + ", ".join(missing)
            )
        
        if status_display:
            status_display.info(
                "🔎 Vérification des colonnes et des références..."
            )

        required = [
            "ID",
            "Description",
            "Marque",
            "Statut",
            "Groupe Produit",
            "Prix (PPC)",
            "Prix (SAV)",
            "Prix HA",
        ]

        source_header_names = {
            normalize(x) for x in get_headers(
                source_ws, int(header_row)
            )
        }

        missing_required = [
            name
            for name in required
            if normalize(name) not in source_header_names
        ]

        if missing_required:
            raise ValueError(
                "Colonnes nécessaires absentes : "
                + ", ".join(missing_required)
            )

        id_col = source_columns[id_column]
        status_col = source_columns[status_column]

        # Dernière ligne : on s'appuie sur l'ID.
        max_row = int(data_start_row) - 1

        for row in range(
            int(data_start_row),
            source_ws.max_row + 1,
        ):
            value = source_ws.cell(
                row=row,
                column=id_col,
            ).value

            if value in (None, ""):
                break

            max_row = row

        total_rows = max(
            0,
            max_row - int(data_start_row) + 1,
        )

        output_wb = openpyxl.Workbook()
        output_ws = output_wb.active
        output_ws.title = sheet_name

        # En-têtes
        for out_col, name in enumerate(output_columns, 1):
            source_cell = source_ws.cell(
                row=int(header_row),
                column=source_columns[name],
            )

            target_cell = output_ws.cell(
                row=1,
                column=out_col,
            )

            target_cell.value = source_cell.value
            copy_style_safe(source_cell, target_cell)

        # Index des IDs existants
        existing_ids = set()

        for row in range(int(data_start_row), max_row + 1):
            value = source_ws.cell(
                row=row,
                column=id_col,
            ).value

            if value not in (None, ""):
                existing_ids.add(normalize(value))

        rows_to_create = []
        vg_count = 0
        pg_count = 0
        duplicate_count = 0
        empty_id_count = 0

        output_row = 2

        def copy_row(source_row, target_row, new_id=None):
            for out_col, name in enumerate(output_columns, 1):
                source_col = source_columns[name]

                source_cell = source_ws.cell(
                    row=source_row,
                    column=source_col,
                )

                target_cell = output_ws.cell(
                    row=target_row,
                    column=out_col,
                )

                # Récupération des valeurs calculées Excel
                # pour les colonnes de prix concernées.
                if normalize(name) in {
                    normalize("Prix (SAV)"),
                    normalize("Prix HA"),
                }:
                    target_cell.value = values_ws.cell(
                        row=source_row,
                        column=source_col,
                    ).value
                else:
                    target_cell.value = source_cell.value

                copy_style_safe(source_cell, target_cell)

            if new_id is not None:
                id_out_col = next(
                    i + 1
                    for i, name in enumerate(output_columns)
                    if normalize(name) == normalize("ID")
                )
                output_ws.cell(
                    row=target_row,
                    column=id_out_col,
                ).value = new_id
        
        if status_display:
            status_display.info(
                "⚙️ Traitement des lignes et détection des statuts VG / PG..."
            )
            
        for index, source_row in enumerate(
            range(int(data_start_row), max_row + 1)
        ):
            copy_row(source_row, output_row)

            status = normalize(
                source_ws.cell(
                    row=source_row,
                    column=status_col,
                ).value
            )

            if status == "VG":
                vg_count += 1
            elif status == "PG":
                pg_count += 1

            original_id = source_ws.cell(
                row=source_row,
                column=id_col,
            ).value

            if status in STATUS_TO_DUPLICATE:
                if original_id in (None, ""):
                    empty_id_count += 1
                else:
                    new_id = f"{str(original_id).strip()}_SAV"
                    normalized_new = normalize(new_id)

                    if normalized_new in existing_ids:
                        duplicate_count += 1
                    else:
                        existing_ids.add(normalized_new)
                        rows_to_create.append((source_row, new_id))

            output_row += 1

            if progress and (
                index % 25 == 0 or index == total_rows - 1
            ):
                pct = 10 + int(
                    75 * (index + 1) / max(1, total_rows)
                )
                progress.progress(
                    min(pct, 85),
                    text=(
                        f"Heidenhain : {index + 1:,} / "
                        f"{total_rows:,}"
                    ),
                )

        for source_row, new_id in rows_to_create:
            copy_row(source_row, output_row, new_id)
            output_row += 1

        # Largeurs de colonnes
        for out_col, name in enumerate(output_columns, 1):
            source_col = source_columns[name]

            source_letter = openpyxl.utils.get_column_letter(
                source_col
            )
            target_letter = openpyxl.utils.get_column_letter(
                out_col
            )

            width = source_ws.column_dimensions[
                source_letter
            ].width

            if width:
                output_ws.column_dimensions[
                    target_letter
                ].width = width

        output_ws.freeze_panes = "A2"

        if status_display:
            status_display.info("💾 Création du fichier Excel...")
            
        result = BytesIO()
        output_wb.save(result)
        result.seek(0)

        stats = {
            "total_rows": total_rows,
            "vg": vg_count,
            "pg": pg_count,
            "created": len(rows_to_create),
            "duplicates": duplicate_count,
            "empty_ids": empty_id_count,
            "total_columns": len(output_columns),
            "columns": list(output_columns),
        }

        output_wb.close()

        if progress:
            progress.progress(100, text="Étape 1 terminée.")

        if status_display:
            status_display.success(
                "Fichier Heidenhain créé."
            )

        return result.getvalue(), stats

    finally:
        wb_formula.close()
        wb_values.close()
        st.session_state.heidenhain_processing = False


# ==========================================================
# CATEGORIES
# ==========================================================

def load_category_mapping(
    category_file,
    category_sheet,
    category_id_col,
    category_search_col,
):
    category_file.seek(0)

    wb = openpyxl.load_workbook(
        category_file,
        read_only=True,
        data_only=True,
    )

    try:
        if category_sheet not in wb.sheetnames:
            raise ValueError(
                f"Feuille catégorie introuvable : {category_sheet}"
            )

        ws = wb[category_sheet]
        mapping = {}
        rows_loaded = 0

        for row in ws.iter_rows(values_only=True):
            max_index = max(
                int(category_id_col),
                int(category_search_col),
            ) - 1

            if max_index >= len(row):
                continue

            category_id = row[int(category_id_col) - 1]
            category_text = row[int(category_search_col) - 1]

            if category_id in (None, ""):
                continue
            if category_text in (None, ""):
                continue

            category_id = str(category_id).strip()
            category_text = str(category_text).strip()

            if not category_id or not category_text:
                continue

            rows_loaded += 1
            mapping.setdefault(
                normalize(category_text),
                category_id,
            )

            for separator in ("-", "–"):
                if separator in category_text:
                    first_part = (
                        category_text.split(separator, 1)[0].strip()
                    )
                    if first_part:
                        mapping.setdefault(
                            normalize(first_part),
                            category_id,
                        )

        return mapping, rows_loaded

    finally:
        wb.close()
        category_file.seek(0)


def find_category_id_fast(value, category_mapping):
    if value in (None, ""):
        return None

    text = str(value).strip()

    result = category_mapping.get(normalize(text))
    if result is not None:
        return result

    for separator in ("-", "–"):
        if separator in text:
            first_part = text.split(separator, 1)[0].strip()
            result = category_mapping.get(normalize(first_part))
            if result is not None:
                return result

    return None


# ==========================================================
# AJOUT DES COLONNES ODOO ABSENTES
# ==========================================================

def ensure_odoo_columns(
    ws,
    header_row,
    required_columns,
    placement_config,
):
    existing_headers = get_headers(ws, header_row)
    existing_normalized = {
        normalize(x) for x in existing_headers
    }

    missing = [
        name
        for name in required_columns
        if normalize(name) not in existing_normalized
    ]

    for name in missing:
        placement = placement_config.get(name, "__END__")

        if placement == "__END__":
            new_col = max(ws.max_column, 1) + 1

            # Si la dernière colonne est vide, on peut
            # réutiliser sa position.
            while new_col > 1 and ws.cell(
                row=header_row,
                column=new_col - 1,
            ).value in (None, ""):
                new_col -= 1

        else:
            target_col = None

            for cell in ws[header_row]:
                if normalize(cell.value) == normalize(placement):
                    target_col = cell.column
                    break

            if target_col is None:
                new_col = max(ws.max_column, 1) + 1
            else:
                ws.insert_cols(target_col, 1)
                new_col = target_col

        # Le style est copié depuis une colonne voisine.
        if new_col > 1:
            copy_style_safe(
                ws.cell(row=header_row, column=new_col - 1),
                ws.cell(row=header_row, column=new_col),
            )

        ws.cell(
            row=header_row,
            column=new_col,
        ).value = name

        existing_normalized.add(normalize(name))


# ==========================================================
# ETAPE 2 : TRAITEMENT ODOO
# ==========================================================

def process_odoo(
    heidenhain_file,
    odoo_file,
    category_file,
    heidenhain_sheet,
    odoo_sheet,
    category_sheet,
    category_id_col,
    category_search_col,
    mapping_config,
    placement_config,
    odoo_header_row,
    heidenhain_output_columns,
    extra_column_names=None,
    progress=None,
    status_display=None,
):
    start_time = time.perf_counter()
    extra_column_names = extra_column_names or {}

    if status_display:
        status_display.info("⚙️ Chargement des catégories...")

    category_mapping, category_rows = load_category_mapping(
        category_file,
        category_sheet,
        category_id_col,
        category_search_col,
    )

    # ------------------------------------------------------
    # Lecture du fichier Heidenhain préparé
    # ------------------------------------------------------

    heidenhain_file.seek(0)

    wb_h = openpyxl.load_workbook(
        heidenhain_file,
        read_only=True,
        data_only=True,
    )

    try:
        if heidenhain_sheet not in wb_h.sheetnames:
            raise ValueError(
                f"Feuille Heidenhain introuvable : {heidenhain_sheet}"
            )

        h_ws = wb_h[heidenhain_sheet]
        h_headers = get_headers(h_ws, 1)
        h_columns = find_columns(h_ws, 1)

        h_columns_normalized = {
            normalize(name): col
            for name, col in h_columns.items()
        }

        def hcol(name):
            return h_columns_normalized.get(normalize(name))

        # Tous les champs réellement utilisés doivent être présents.
        required_sources = set(REQUIRED_HEIDENHAIN_COLUMNS)

        for config in mapping_config.values():
            source = config.get("source")
            if source and not source.startswith("__"):
                required_sources.add(source)

        missing_sources = [
            name
            for name in required_sources
            if hcol(name) is None
        ]

        if missing_sources:
            raise ValueError(
                "Colonnes nécessaires absentes du fichier préparé "
                "à l'étape 1 : "
                + ", ".join(sorted(missing_sources))
                + ". Vérifie la sélection des colonnes."
            )

        missing_selected = [
            name
            for name in heidenhain_output_columns
            if hcol(name) is None
        ]

        if missing_selected:
            raise ValueError(
                "Colonnes sélectionnées à l'étape 1 introuvables : "
                + ", ".join(missing_selected)
            )

        # --------------------------------------------------
        # Ouverture du fichier Odoo
        # --------------------------------------------------

        odoo_file.seek(0)

        wb_o = openpyxl.load_workbook(
            odoo_file,
            data_only=False,
        )

        try:
            if odoo_sheet not in wb_o.sheetnames:
                raise ValueError(
                    f"Feuille Odoo introuvable : {odoo_sheet}"
                )

            ws = wb_o[odoo_sheet]
            header_row = int(odoo_header_row)

            original_headers = get_headers(ws, header_row)
            original_normalized = {
                normalize(x) for x in original_headers
            }

            # Détection de l'ID avant d'ajouter les champs manquants.
            id_exists = normalize("id") in original_normalized

            required_odoo_columns = build_required_odoo_columns(
                heidenhain_output_columns,
                mapping_config,
            )
            
            # Ne jamais ajouter la colonne ID si elle n'existe pas
            # déjà dans le fichier modèle Odoo.
            if not id_exists:
                required_odoo_columns = [
                    col for col in required_odoo_columns
                    if normalize(col) != normalize("id")
                ]
            
            for source_name, target_name in extra_column_names.items():
                if not target_name.strip():
                    raise ValueError(
                        f"Le nom Odoo de la colonne "
                        f"« {source_name} » est obligatoire."
                    )

                required_odoo_columns = [
                    col
                    for col in required_odoo_columns
                    if normalize(col) != normalize(source_name)
                ]

                if normalize(target_name) not in {
                    normalize(col)
                    for col in required_odoo_columns
                }:
                    required_odoo_columns.append(target_name)

            ensure_odoo_columns(
                ws,
                header_row,
                required_odoo_columns,
                placement_config,
            )

            odoo_columns = find_columns(
                ws,
                header_row,
                required_odoo_columns,
            )

            missing_odoo = [
                name
                for name in required_odoo_columns
                if name not in odoo_columns
            ]

            if missing_odoo:
                raise ValueError(
                    "Impossible de créer les colonnes Odoo : "
                    + ", ".join(missing_odoo)
                )

            # --------------------------------------------------
            # Index des références Odoo
            # --------------------------------------------------

            reference_col = odoo_columns["default_code"]
            reference_index = {}

            for row in range(header_row + 1, ws.max_row + 1):
                value = ws.cell(
                    row=row,
                    column=reference_col,
                ).value

                ref = clean_reference(value)

                if ref:
                    reference_index.setdefault(ref, row)

            
            # Première ligne réellement disponible après les données existantes.
            # Ignore les lignes contenant uniquement du formatage Excel.
            last_data_row = header_row
            
            for row in range(header_row + 1, ws.max_row + 1):
                if any(
                    ws.cell(row=row, column=col).value not in (None, "")
                    for col in range(1, ws.max_column + 1)
                ):
                    last_data_row = row
            
            next_empty_row = last_data_row + 1

            # --------------------------------------------------
            # Détermination des colonnes supplémentaires
            # --------------------------------------------------

            mapped_sources = {
                normalize(config.get("source"))
                for config in mapping_config.values()
                if config.get("source")
                and not config.get("source").startswith("__")
            }

            extra_columns = [
                name
                for name in heidenhain_output_columns
                if normalize(name) not in mapped_sources
                and normalize(name) not in {
                    normalize("Prix (PPC)"),
                    normalize("Prix (SAV)"),
                }
            ]

            processed = 0
            updated = 0
            created = 0
            category_found = 0
            category_missing = 0
            empty_references = 0

            total_rows = max(0, h_ws.max_row - 1)

            # --------------------------------------------------
            # Traitement des lignes
            # --------------------------------------------------

            for index, row_values in enumerate(
                h_ws.iter_rows(
                    min_row=2,
                    values_only=True,
                )
            ):
                id_index = hcol("ID") - 1
                id_value = row_values[id_index]

                if id_value in (None, ""):
                    break

                reference = str(id_value).strip()

                if not reference:
                    empty_references += 1
                    continue

                processed += 1
                normalized_ref = clean_reference(reference)

                if normalized_ref in reference_index:
                    target_row = reference_index[normalized_ref]
                    updated += 1
                else:
                    target_row = next_empty_row
                    next_empty_row += 1

                    reference_index[normalized_ref] = target_row
                    created += 1

                def source_value(name):
                    source_col = hcol(name)

                    if source_col is None:
                        return None

                    return row_values[source_col - 1]

                status = source_value("Statut")
                groupe = source_value("Groupe Produit")
                ppc = source_value("Prix (PPC)")
                sav = source_value("Prix (SAV)")
                prix_ha = source_value("Prix HA")

                price = sav if is_sav(reference) else ppc

                category_id = find_category_id_fast(
                    groupe,
                    category_mapping,
                )

                if category_id is None:
                    category_missing += 1
                else:
                    category_found += 1

                # --------------------------------------------------
                # Correspondances Odoo
                # --------------------------------------------------

                for field, config in mapping_config.items():
                    target_col = odoo_columns[field]

                    source = config.get("source")
                    fixed = config.get("fixed")

                    # Ne jamais écraser un ID Odoo.
                    if normalize(field) == normalize("id"):
                        continue

                    if source == "__PRICE__":
                        value = price

                    elif source == "__BARCODE__":
                        value = (
                            ""
                            if is_sav(reference)
                            else f"I {reference}"
                        )

                    elif source is not None:
                        value = source_value(source)

                        if normalize(field) == normalize("categ_id"):
                            # Ne pas effacer une catégorie existante
                            # si aucune correspondance n'est trouvée.
                            if category_id is None:
                                continue
                            value = category_id

                    else:
                        value = fixed

                    ws.cell(
                        row=target_row,
                        column=target_col,
                    ).value = value

                # --------------------------------------------------
                # Transfert des colonnes supplémentaires
                # --------------------------------------------------

                for source_name in extra_columns:
                    source_col = hcol(source_name)

                    if source_col is None:
                        continue

                    # Utiliser le nom Odoo configuré par l'utilisateur.
                    target_name = extra_column_names.get(
                        source_name,
                        source_name,
                    )

                    target_col = odoo_columns.get(target_name)

                    if target_col is None:
                        continue

                    ws.cell(
                        row=target_row,
                        column=target_col,
                    ).value = row_values[source_col - 1]

                # --------------------------------------------------
                # Progression
                # --------------------------------------------------

                if progress and (
                    index % 25 == 0
                    or index == total_rows - 1
                ):
                    pct = 10 + int(
                        80 * (index + 1) / max(1, total_rows)
                    )

                    progress.progress(
                        min(pct, 90),
                        text=(
                            f"Odoo : {index + 1:,} / {total_rows:,}"
                            f" | Mises à jour : {updated:,}"
                            f" | Créations : {created:,}"
                        ),
                    )

            # --------------------------------------------------
            # Sauvegarde
            # --------------------------------------------------

            if status_display:
                status_display.info("💾 Création du fichier Odoo final...")

            output = BytesIO()
            wb_o.save(output)
            output.seek(0)

            stats = {
                "processed": processed,
                "updated": updated,
                "created": created,
                "category_found": category_found,
                "category_missing": category_missing,
                "empty_references": empty_references,
                "category_rows": category_rows,
                "odoo_indexed": len(reference_index),
                "id_exists": id_exists,
                "total_columns": len(
                    get_headers(ws, header_row)
                ),
                "total_time": time.perf_counter() - start_time,
            }

            if progress:
                progress.progress(100, text="Étape 2 terminée.")

            if status_display:
                status_display.success(
                    "Fichier Odoo préparé avec succès."
                )

            return output.getvalue(), stats

        finally:
            wb_o.close()

    finally:
        wb_h.close()
        st.session_state.Odoo_processing = False


# ==========================================================
# INTERFACE : FICHIERS
# ==========================================================

st.header("📂 Fichiers")

col1, col2, col3 = st.columns(3)

with col1:
    st.subheader("📘 Prix Heidenhain")
    heidenhain_file = st.file_uploader(
        "Fichier Heidenhain",
        type=["xlsx", "xlsm"],
        key="heidenhain_file",
    )

with col2:
    st.subheader("📗 Fichier Odoo")
    odoo_file = st.file_uploader(
        "Fichier d'import Odoo",
        type=["xlsx", "xlsm"],
        key="odoo_file",
    )

with col3:
    st.subheader("📂 Catégories")
    category_file = st.file_uploader(
        "Fichier catégories",
        type=["xlsx", "xlsm"],
        key="category_file",
    )


# ==========================================================
# PARAMETRES
# ==========================================================

st.sidebar.header("⚙️ Paramètres")

# ----------------------------------------------------------
# Heidenhain
# ----------------------------------------------------------

st.sidebar.subheader("📘 Fichier Heidenhain")

heidenhain_sheet = st.sidebar.text_input(
    "Feuille Heidenhain",
    value="Distributeurs",
)

heidenhain_header_row = st.sidebar.number_input(
    "Ligne des en-têtes Heidenhain",
    min_value=1,
    value=4,
    step=1,
)

heidenhain_detected_columns = []

if heidenhain_file is not None:
    heidenhain_detected_columns = detect_excel_headers(
        heidenhain_file,
        heidenhain_sheet,
        heidenhain_header_row,
    )

    if heidenhain_detected_columns:
        st.sidebar.success(
            f"{len(heidenhain_detected_columns)} "
            "colonnes Heidenhain détectées."
        )
    else:
        st.sidebar.warning(
            "Aucune colonne détectée. Vérifie la feuille "
            "et la ligne des en-têtes."
        )

st.sidebar.subheader("📋 Colonnes à conserver à l'étape 1")

if heidenhain_detected_columns:
    default_columns = [
        name
        for name in DEFAULT_HEIDENHAIN_COLUMNS
        if normalize(name) in {
            normalize(x) for x in heidenhain_detected_columns
        }
    ]

    heidenhain_output_columns = st.sidebar.multiselect(
        "Sélectionner les colonnes à garder",
        options=heidenhain_detected_columns,
        default=default_columns,
        help=(
            "Seules les colonnes prévues sont cochées par défaut. "
            "Les autres colonnes restent sélectionnables."
        ),
        key="heidenhain_output_columns",
    )
else:
    heidenhain_output_columns = []

# ----------------------------------------------------------
# Odoo
# ----------------------------------------------------------

st.sidebar.divider()
st.sidebar.subheader("📗 Fichier Odoo")

odoo_sheet = st.sidebar.text_input(
    "Feuille Odoo",
    value="Sheet1",
)

odoo_header_row = st.sidebar.number_input(
    "Ligne des en-têtes Odoo",
    min_value=1,
    value=1,
    step=1,
)

odoo_detected_columns = []

if odoo_file is not None:
    odoo_detected_columns = detect_excel_headers(
        odoo_file,
        odoo_sheet,
        odoo_header_row,
    )

    if odoo_detected_columns:
        st.sidebar.success(
            f"{len(odoo_detected_columns)} "
            "colonnes Odoo détectées."
        )

        if normalize("id") in {
            normalize(x) for x in odoo_detected_columns
        }:
            st.sidebar.info(
                "Mode import Odoo activé : "
                "l'ID Odoo existant est conservé."
            )
    else:
        st.sidebar.warning(
            "Aucune colonne Odoo détectée."
        )


# ==========================================================
# PARAMETRES CATEGORIES
# ==========================================================

st.sidebar.divider()
st.sidebar.subheader("📂 Catégories")

category_sheet = st.sidebar.text_input(
    "Feuille catégories",
    value="Sheet1",
)

category_id_col = st.sidebar.number_input(
    "Numéro de colonne ID catégorie",
    min_value=1,
    value=2,
    step=1,
)

category_search_col = st.sidebar.number_input(
    "Numéro de colonne nom catégorie",
    min_value=1,
    value=3,
    step=1,
)

# ==========================================================
# CORRESPONDANCES ODOO
# Renommage des colonnes Heidenhain supplémentaires
# ==========================================================

placement_config = {}
extra_column_names = {}

st.sidebar.subheader("🔗 Correspondances Odoo")

if odoo_file is not None and odoo_detected_columns:

    odoo_normalized = {
        normalize(col): col
        for col in odoo_detected_columns
    }

    mapped_sources = {
        normalize(config["source"])
        for config in ODOO_FIELD_CONFIGURATION.values()
        if config.get("source")
        and not config["source"].startswith("__")
    }

    missing_standard = [
        field
        for field in ODOO_FIELD_CONFIGURATION
        if normalize(field) not in odoo_normalized
    ]

    extra_columns = [
        col
        for col in heidenhain_output_columns
        if normalize(col) not in mapped_sources
        and normalize(col) not in {
            normalize("Prix (PPC)"),
            normalize("Prix (SAV)"),
        }
    ]

    missing_extra = [
        col
        for col in extra_columns
        if normalize(col) not in odoo_normalized
    ]

    # Champs standards : noms techniques prédéfinis
    for field in missing_standard:
        placement_config[field] = "__END__"

    # Colonnes supplémentaires : demander le nom Odoo
    if missing_extra:
        st.sidebar.warning(
            f"{len(missing_extra)} colonne(s) supplémentaire(s) "
            "à configurer."
        )

        st.sidebar.caption(
            "Indique le nom de chaque colonne à créer dans le fichier "
            "Odoo. Le nom proposé par défaut est celui de Heidenhain."
        )

        for source_name in missing_extra:
            target_name = st.sidebar.text_input(
                f"Nom de la colonne Odoo pour « {source_name} »",
                value=source_name,
                key=f"odoo_target_name_{normalize(source_name)}",
                help=(
                    "Exemple : ROHS → x_studio_rohs. "
                    "Le nom sera utilisé comme en-tête dans le fichier Excel."
                ),
            ).strip()

            if target_name:
                extra_column_names[source_name] = target_name
                placement_config[target_name] = "__END__"

    if not missing_extra:
        st.sidebar.success(
            "✅ Toutes les correspondances Odoo sont configurées."
        )

        with st.sidebar.expander(
            "📋 Voir les correspondances actives",
            expanded=False,
        ):
            st.markdown("**Champs standards**")

            for field, config in ODOO_FIELD_CONFIGURATION.items():
                source = config.get("source")
                fixed = config.get("fixed")

                if source == "__PRICE__":
                    origine = "Prix PPC / SAV"
                elif source == "__BARCODE__":
                    origine = "Code-barres calculé"
                elif source:
                    origine = source
                elif fixed is not None:
                    origine = f"Valeur fixe : {fixed}"
                else:
                    origine = "ID Odoo conservé"

                st.write(f"**{field}** ← {origine}")

            if extra_column_names:
                st.markdown("**Colonnes supplémentaires**")

                for source_name, target_name in extra_column_names.items():
                    st.write(f"**{target_name}** ← {source_name}")

# ==========================================================
# ETAPE 1
# ==========================================================

st.divider()
st.header("1️⃣ Création du fichier Heidenhain")

if heidenhain_file is None:
    st.info("Importe le fichier Heidenhain pour commencer.")
else:
    st.write(
        "Les colonnes cochées par défaut sont : "
        + ", ".join(
            name
            for name in DEFAULT_HEIDENHAIN_COLUMNS
            if normalize(name) in {
                normalize(x) for x in heidenhain_output_columns
            }
        )
    )

    if not st.session_state.heidenhain_processing:
        if st.button(
            "🚀 Créer le fichier Heidenhain",
            type="primary",
            use_container_width=True,
            key="create_heidenhain",
        ):
            st.session_state.heidenhain_processing = True
            st.rerun()
    
    else:
        st.button(
            "⏳ Traitement Heidenhain en cours...",
            disabled=True,
            use_container_width=True,
            key="create_heidenhain_disabled",
        )
    
    if st.session_state.heidenhain_processing:
        progress = st.progress(0, text="Initialisation...")
        status = st.empty()
    
        try:
            result, stats = process_heidenhain(
                uploaded_file=heidenhain_file,
                sheet_name=heidenhain_sheet,
                header_row=int(heidenhain_header_row),
                data_start_row=int(heidenhain_header_row) + 1,
                status_column="Statut",
                id_column="ID",
                output_columns=heidenhain_output_columns,
                progress=progress,
                status_display=status,
            )
    
            st.session_state.heidenhain_result = result
            st.session_state.heidenhain_stats = stats
            st.session_state.odoo_result = None
            st.session_state.odoo_stats = None
    
        except Exception as exc:
            st.error(f"Erreur pendant la préparation Heidenhain : {exc}")
            st.exception(exc)
    
        finally:
            st.session_state.heidenhain_processing = False
            st.rerun()

if st.session_state.heidenhain_result is not None:
    stats = st.session_state.heidenhain_stats or {}

    total_columns = stats.get(
        "total_columns",
        len(stats.get("columns", [])),
    )

    st.success(
        f"Fichier créé : {total_columns} colonnes conservées, "
        f"{stats.get('created', 0)} lignes SAV ajoutées."
    )

    st.download_button(
        "⬇️ Télécharger le fichier Heidenhain",
        data=st.session_state.heidenhain_result,
        file_name="ETAPE 1 - Products_Heidenhain_Prepare.xlsx",
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        use_container_width=True,
        key="download_heidenhain",
    )

    with st.expander("Colonnes présentes dans le fichier créé"):
        st.write(stats.get("columns", []))


# ==========================================================
# ETAPE 2
# ==========================================================

st.divider()
st.header("2️⃣ Préparation du fichier Odoo")

if (
    st.session_state.heidenhain_result is None
    or odoo_file is None
    or category_file is None
):
    st.info(
        "Crée d'abord le fichier Heidenhain, puis importe "
        "le fichier Odoo et le fichier catégories."
    )
else:
    if not st.session_state.odoo_processing:
        if st.button(
            "🚀 Préparer le fichier Odoo",
            type="primary",
            use_container_width=True,
            key="prepare_odoo",
        ):
            st.session_state.odoo_processing = True
            st.rerun()

    else:
        st.button(
            "⏳ Traitement Odoo en cours...",
            disabled=True,
            use_container_width=True,
            key="prepare_odoo_disabled",
        )

    if st.session_state.odoo_processing:
        progress = st.progress(0, text="Initialisation...")
        status = st.empty()

        try:
            prepared_heidenhain = BytesIO(
                st.session_state.heidenhain_result
            )

            result, stats = process_odoo(
                heidenhain_file=prepared_heidenhain,
                odoo_file=odoo_file,
                category_file=category_file,
                heidenhain_sheet=heidenhain_sheet,
                odoo_sheet=odoo_sheet,
                category_sheet=category_sheet,
                category_id_col=int(category_id_col),
                category_search_col=int(category_search_col),
                mapping_config=ODOO_FIELD_CONFIGURATION,
                placement_config=placement_config,
                odoo_header_row=int(odoo_header_row),
                heidenhain_output_columns=heidenhain_output_columns,
                extra_column_names=extra_column_names,
                progress=progress,
                status_display=status,
            )

            st.session_state.odoo_result = result
            st.session_state.odoo_stats = stats

        except Exception as exc:
            st.error(
                "Erreur pendant la préparation Odoo. "
                "Vérifie les colonnes et leurs emplacements."
            )
            st.exception(exc)

        finally:
            st.session_state.odoo_processing = False
            st.rerun()


# ==========================================================
# RESULTAT ODOO
# ==========================================================

if st.session_state.odoo_result is not None:
    stats = st.session_state.odoo_stats

    st.divider()
    st.header("📊 Résultat Odoo")

    c1, c2, c3, c4, c5 = st.columns(5)

    c1.metric("Références traitées", f"{stats['processed']:,}")
    c2.metric("Lignes mises à jour", f"{stats['updated']:,}")
    c3.metric("Nouvelles lignes", f"{stats['created']:,}")
    c4.metric("Catégories trouvées", f"{stats['category_found']:,}")
    c5.metric("Catégories absentes", f"{stats['category_missing']:,}")

    st.info(
        f"{stats['total_columns']} colonnes dans le fichier Odoo final."
    )

    if stats["id_exists"]:
        st.info(
            "Mode import Odoo activé : l'ID Odoo existant est conservé."
        )

    if stats["category_missing"]:
        st.warning(
            f"{stats['category_missing']:,} référence(s) "
            "sans correspondance de catégorie."
        )

    st.info(
        f"{stats['category_rows']:,} lignes de catégories chargées."
    )

    st.info(
        f"Temps total : {stats['total_time']:.2f} secondes."
    )

    st.download_button(
        "⬇️ Télécharger le fichier Odoo prêt à importer",
        data=st.session_state.odoo_result,
        file_name="ETAPE 2 - File for import Odoo.xlsx",
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        use_container_width=True,
        key="download_odoo",
    )

    st.success("Le fichier Odoo final est prêt.")
